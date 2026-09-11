"""
Generates deterministic synthetic engineering test documents containing realistic deviations.
Adheres strictly to the user prompt test requirements:
1. Mechanical: Expected tolerance ±0.05 mm vs Detected ±0.5 mm
2. Electrical: Expected voltage 415 V vs Detected 230 V
3. Chemical: Expected concentration 10% vs Detected 15%
4. Logical: Page 1 Max temp 80°C vs Page 2 Max temp 60°C
5. Structural: 3.1 -> 3.3 (missing 3.2)
6. TOC: TOC page drift
7. Grammar: Controlled spelling errors & repeated words
8. Tables: Missing cells and unit mixing
"""

import os
from pathlib import Path
import pymupdf
import docx

DEMO_DIR = Path(__file__).resolve().parent / "demo_samples"
DEMO_DIR.mkdir(parents=True, exist_ok=True)


def create_mechanical_sample():
    doc = pymupdf.open()

    # PAGE 1: Title, TOC, Section 1 & 2
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text((54, 60), "SPECIFICATION FOR PRECISION SHAFT & FLANGE ASSEMBLY", fontsize=15, fontname="helv", color=(0.1, 0.2, 0.4))
    p1.insert_text((54, 85), "Document No: MEC-ENG-2024-001 | Revision: C", fontsize=10, fontname="helv", color=(0.4, 0.4, 0.4))

    # TOC with deliberate page number drift for Section 3 (lists Page 4, but Section 3 is on Page 2)
    p1.insert_text((54, 120), "Table of Contents", fontsize=12, fontname="helv", color=(0.2, 0.2, 0.2))
    p1.insert_text((54, 140), "1 Scope ........................................................... 1", fontsize=10, fontname="helv")
    p1.insert_text((54, 158), "2 System Requirements ................................... 1", fontsize=10, fontname="helv")
    p1.insert_text((54, 176), "3 Manufacturing Tolerances ............................ 4", fontsize=10, fontname="helv") # DELIBERATE TOC ERROR: Actually on Page 2

    # Section 1
    p1.insert_text((54, 220), "1 Scope", fontsize=13, fontname="helv", color=(0.1, 0.2, 0.4))
    p1.insert_text((54, 240), "This engineering specification governs the precision machining, dimensional limits,", fontsize=10, fontname="helv")
    p1.insert_text((54, 255), "and quality inspection for heavy industrial centrifugal pump drive shafts.", fontsize=10, fontname="helv")

    # Section 2 with initial temperature and grammar error
    p1.insert_text((54, 290), "2 System Requirements", fontsize=13, fontname="helv", color=(0.1, 0.2, 0.4))
    p1.insert_text((54, 310), "The pump shaft operates in harsh environmental conditions.", fontsize=10, fontname="helv")
    p1.insert_text((54, 325), "The maximum operating temperature = 80°C under continuous rated load.", fontsize=10, fontname="helv") # First temp
    p1.insert_text((54, 340), "Verify that all maintanence schedules are strictly followed.", fontsize=10, fontname="helv") # DELIBERATE SPELLING: maintanence
    p1.insert_text((54, 355), "The assembly must withstand maximum operating pressure = 10 bar without leakage.", fontsize=10, fontname="helv")

    # PAGE 2: Section 3 with broken numbering (3.1 -> 3.3), tolerance error, and logical temp contradiction
    p2 = doc.new_page(width=612, height=792)
    p2.insert_text((54, 60), "3 Manufacturing Tolerances", fontsize=13, fontname="helv", color=(0.1, 0.2, 0.4))

    p2.insert_text((54, 90), "3.1 Shaft Machining Limits", fontsize=12, fontname="helv", color=(0.2, 0.2, 0.3))
    # DELIBERATE CRITICAL TOLERANCE ERROR: ±0.5 mm instead of standard ±0.05 mm
    p2.insert_text((54, 110), "Critical bearing journal diameter tolerance = ±0.5 mm on all bearing seating zones.", fontsize=10, fontname="helv")
    p2.insert_text((54, 125), "The specified surface roughness = 3.2 um for general unground surfaces.", fontsize=10, fontname="helv")

    # DELIBERATE STRUCTURAL SEQUENCE ERROR: 3.3 follows 3.1 directly (Missing 3.2)
    p2.insert_text((54, 170), "3.3 Environmental Operating Limits", fontsize=12, fontname="helv", color=(0.2, 0.2, 0.3))
    # DELIBERATE LOGICAL CONTRADICTION: Maximum temperature = 60°C (conflicts with 80°C on Page 1)
    p2.insert_text((54, 190), "The drive shaft maximum operating temperature = 60°C during continuous operation.", fontsize=10, fontname="helv")
    p2.insert_text((54, 205), "Note that teh pump casing was recieved in good condition.", fontsize=10, fontname="helv") # DELIBERATE TYPO: teh, recieved

    # Draw table with missing cell and unit inconsistency
    p2.insert_text((54, 250), "Table 1: Critical Component Dimensions", fontsize=11, fontname="helv", color=(0.1, 0.2, 0.4))
    # Table headers
    p2.draw_rect(pymupdf.Rect(54, 265, 540, 365), color=(0.3, 0.3, 0.3), width=1)
    p2.draw_line((54, 285), (540, 285), color=(0.3, 0.3, 0.3), width=1)
    p2.draw_line((200, 265), (200, 365), color=(0.3, 0.3, 0.3), width=1)
    p2.draw_line((360, 265), (360, 365), color=(0.3, 0.3, 0.3), width=1)

    p2.insert_text((60, 280), "Component", fontsize=10, fontname="helv")
    p2.insert_text((210, 280), "Nominal Dimension", fontsize=10, fontname="helv")
    p2.insert_text((370, 280), "Machining Tolerance", fontsize=10, fontname="helv")

    # Row 1
    p2.insert_text((60, 305), "Shaft Journal", fontsize=9, fontname="helv")
    p2.insert_text((210, 305), "50 mm", fontsize=9, fontname="helv")
    p2.insert_text((370, 305), "±0.05 mm", fontsize=9, fontname="helv")

    # Row 2 (DELIBERATE EMPTY CELL in critical tolerance)
    p2.insert_text((60, 330), "Flange Collar", fontsize=9, fontname="helv")
    p2.insert_text((210, 330), "120 mm", fontsize=9, fontname="helv")
    p2.insert_text((370, 330), "-", fontsize=9, fontname="helv") # EMPTY CELL

    # Row 3 (DELIBERATE MIXED UNIT in column 2: 12.5 cm instead of mm)
    p2.insert_text((60, 355), "Coupling Hub", fontsize=9, fontname="helv")
    p2.insert_text((210, 355), "12.5 cm", fontsize=9, fontname="helv") # INCONSISTENT UNIT
    p2.insert_text((370, 355), "±0.02 mm", fontsize=9, fontname="helv")

    out_file = DEMO_DIR / "mechanical_sample_with_errors.pdf"
    doc.save(str(out_file))
    doc.close()
    print(f"Created {out_file}")


def create_electrical_sample():
    doc = pymupdf.open()
    p1 = doc.new_page(width=612, height=792)
    p1.insert_text((54, 60), "ELECTRICAL MOTOR CONTROL CENTER SPECIFICATION", fontsize=15, fontname="helv", color=(0.1, 0.2, 0.4))
    p1.insert_text((54, 85), "Project: Industrial Plant MCC-01 | Standard: IEC 60364", fontsize=10, fontname="helv", color=(0.4, 0.4, 0.4))

    p1.insert_text((54, 130), "1 Scope", fontsize=13, fontname="helv", color=(0.1, 0.2, 0.4))
    p1.insert_text((54, 150), "This specification defines the electrical distribution and motor feeder ratings.", fontsize=10, fontname="helv")

    p1.insert_text((54, 190), "2 System Requirements", fontsize=13, fontname="helv", color=(0.1, 0.2, 0.4))
    # DELIBERATE CRITICAL VOLTAGE DEVIATION: 230 V instead of expected 415 V for three-phase motor feeder
    p1.insert_text((54, 210), "The three-phase induction motor operates at a rated voltage = 230 V.", fontsize=10, fontname="helv")
    p1.insert_text((54, 225), "System line operating frequency = 50 Hz with symmetrical sinusoidal supply.", fontsize=10, fontname="helv")
    # DELIBERATE CRITICAL CURRENT OVERLOAD: 750 A exceeds 630 A standard frame limit
    p1.insert_text((54, 240), "Main busbar continuous incoming current = 750 A under peak operational loading.", fontsize=10, fontname="helv")
    # DELIBERATE REPEATED WORD
    p1.insert_text((54, 255), "Check that the the circuit breaker trip unit is correctly calibrated.", fontsize=10, fontname="helv")

    out_file = DEMO_DIR / "electrical_sample_with_errors.pdf"
    doc.save(str(out_file))
    doc.close()
    print(f"Created {out_file}")


def create_chemical_sample():
    doc = docx.Document()
    doc.add_heading("Chemical Process Reactor Operating Specification", level=1)
    doc.add_paragraph("Doc Ref: CPR-SPEC-2024-B | Classification: Safety Critical")

    doc.add_heading("1 Scope", level=2)
    doc.add_paragraph("This document specifies the operational concentration limits and thermal thresholds for batch reactor RX-101.")

    doc.add_heading("2 Requirements", level=2)
    # DELIBERATE CRITICAL CONCENTRATION DEVIATION: 15% exceeds 10% safety standard
    doc.add_paragraph("The aqueous reactant feed required concentration = 15% during steady-state synthesis.")
    # DELIBERATE TEMPERATURE RUNAWAY VIOLATION: 140°C exceeds 120°C standard limit
    doc.add_paragraph("Maximum operating temperature = 140°C in the core reactor vessel jacket.")
    doc.add_paragraph("Ensure that the the emergency quench system valve is primed.") # Repeated word

    out_file = DEMO_DIR / "chemical_sample_with_errors.docx"
    doc.save(str(out_file))
    print(f"Created {out_file}")


if __name__ == "__main__":
    create_mechanical_sample()
    create_electrical_sample()
    create_chemical_sample()
